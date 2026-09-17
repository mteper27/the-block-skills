# -*- coding: utf-8 -*-
import re, sys
sys.path.insert(0,'.'); import content as C
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x1F,0x38,0x64); GREY=RGBColor(0x59,0x59,0x59); RED=RGBColor(0xC0,0,0)
doc=Document()
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin=s.right_margin=Inches(0.55); s.top_margin=s.bottom_margin=Inches(0.5)
st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')

def runs(par, text, size=None, color=None, bold=None, italic=None):
    # inline markdown: **bold**, *italic*, `code`
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not tok: continue
        r=par.add_run()
        if tok.startswith('**'): r.text=tok[2:-2]; r.bold=True
        elif tok.startswith('`'): r.text=tok[1:-1]; r.font.name='Consolas'
        elif tok.startswith('*'): r.text=tok[1:-1]; r.italic=True
        else: r.text=tok
        if size: r.font.size=Pt(size)
        if color is not None: r.font.color.rgb=color
        if bold: r.bold=True
        if italic: r.italic=True
    return par

def shade(cell, hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)

def header(kind):
    p=doc.add_paragraph(); runs(p, C.TITLE, size=17, color=NAVY, bold=True); p.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(); runs(p, C.SUB, size=10, color=NAVY); p.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(); runs(p, C.BYLINE+'   ·   ', size=8.5, color=GREY); runs(p, kind, size=8.5, color=GREY, bold=True)
    p.paragraph_format.space_after=Pt(8)

def table(headers, rows, widths, hdr_size=8.2, body_size=8.4):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,hname in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'1F3864'); c.width=widths[i]
        p=c.paragraphs[0]; runs(p,hname,size=hdr_size,color=RGBColor(255,255,255),bold=True)
    for row in rows:
        cells=t.add_row().cells
        for i,(txt,style) in enumerate(row):
            cells[i].width=widths[i]; p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
            runs(p, txt, size=body_size, color=(NAVY if style=='ref' else (GREY if style in('lab','who') else None)),
                 bold=(style in('ref','n','lever')), italic=(style=='lab'))
    sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(0); sp.paragraph_format.space_before=Pt(0)
    for r in sp.runs: r.font.size=Pt(2)
    return t

def h3(text, color=NAVY):
    p=doc.add_paragraph(); runs(p, text.upper(), size=9.6, color=color, bold=True)
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(1)

def pagebreak(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# PAGE 1 — one table, section divider rows, so it stays on one page
header('1 · QUESTIONS'); p=doc.add_paragraph(); runs(p, C.Q_INTRO, size=8.8, color=GREY); p.paragraph_format.space_after=Pt(3)
W=[Inches(0.28),Inches(1.05),Inches(2.35),Inches(3.7)]
t=doc.add_table(rows=1, cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,hname in enumerate(['','Cell / line','Your line reads','Question']):
    c=t.rows[0].cells[i]; shade(c,'1F3864'); c.width=W[i]; p=c.paragraphs[0]
    p.paragraph_format.space_after=Pt(0); runs(p,hname,size=8.2,color=RGBColor(255,255,255),bold=True)
n=0
for sheet, qs in C.QUESTIONS:
    row=t.add_row(); m=row.cells[0].merge(row.cells[3]); shade(m,'D9E2F3'); p=m.paragraphs[0]
    p.paragraph_format.space_after=Pt(0); runs(p, sheet.upper(), size=8.2, color=NAVY, bold=True)
    for ref,lab,q in qs:
        n+=1; cells=t.add_row().cells
        for i,(txt,style) in enumerate([(str(n),'n'),(ref,'ref'),(lab,'lab'),(q,None)]):
            cells[i].width=W[i]; p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
            runs(p, txt, size=8.3, color=(NAVY if style in('ref','n') else (GREY if style=='lab' else None)),
                 bold=(style in('ref','n')), italic=(style=='lab'))
pagebreak()

# PAGE 2
header('2 · RECOMMENDATIONS — WHAT WE CAN DO TO CUT COSTS'); p=doc.add_paragraph(); runs(p, C.R_INTRO, size=9, color=GREY)
W2=[Inches(0.3),Inches(1.7),Inches(1.0),Inches(4.2)]
table(['','Lever','Who','What and why'], [[(str(i),'n'),(l,'lever'),(w,'who'),(x,None)] for i,(l,w,x) in enumerate(C.RECS,1)], W2)
h3('Two things we do not want cut', RED)
for k,v in C.R_DONT:
    p=doc.add_paragraph(); runs(p, f'**{k}.** {v}', size=9)
pagebreak()

# CLARIFICATION
header('3 · CLARIFICATION — THE REASONING BEHIND PAGES 1 AND 2')
for title, paras in C.CLAR:
    p=doc.add_paragraph(); runs(p, title, size=12.5, color=NAVY, bold=True); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(3)
    for para in paras:
        if para.strip().startswith('|'):
            rows=[r for r in para.strip().split('\n') if r.strip() and not re.match(r'^\|\s*-',r)]
            cells=[[c.strip() for c in r.strip('|').split('|')] for r in rows]
            table(cells[0], [[(c,None) for c in r] for r in cells[1:]], [Inches(3.6),Inches(1.8),Inches(1.8)])
        else:
            p=doc.add_paragraph(); runs(p, para, size=9.5); p.paragraph_format.space_after=Pt(5)

OUT='165_Randolph_The_BLOCK_Questions_Recommendations_Clarification.docx'
doc.save(OUT); print('docx ok', OUT)
