# -*- coding: utf-8 -*-
import html, subprocess, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
E='—'; Q='’'
TITLE='165 Randolph Ave — “The BLOCK”'
HEAD=[('To','Thomas'),('cc','Nate (Wake)'),('From','Matthew Teper'),('Date','17 September 2026'),
      ('Re',f'Questions and recommendations on the 16 September Master Development Budget')]
PARAS=[
 "Thomas,",
 f"Thank you for the 16 September revision. The staffing chart behind general conditions, the newly priced electrical, plumbing and finishes, and the benchmark notes you left in the margins are a real step forward from the 4 September file, and we have leaned on all three.",
 f"Attached is one document in three parts. Page 1 is a single page of questions, keyed to your cells and line numbers so you can work through them in your own file. Page 2 is a single page of recommendations {E} what we think can come out, in the order it is worth doing, with who owns each. The pages behind those are the reasoning, for whoever wants it. You should not need them to answer the first two.",
 f"Two things up front. First, we got two things wrong in our earlier read {E} the roof-demolition line and the 135,953 SF basis {E} and we have withdrawn both. We would rather say that plainly than pretend otherwise. Second, the biggest question in the pack is structural rather than arithmetic. The RFEM tonnage prices a full rebuild. Our scheme keeps the existing roof and lifts it, with the stage, the production mezzanine and the VIP mezzanine standing on their own footings. We are not asking for a re-price; we are asking whether the lift is possible, starting with what the existing roof is made of. That one is for Nate and the structural engineer as much as for you.",
 f"If you can send back the five questions at the top of page 1, division 055100 priced off the takeoff, and your broker{Q}s quote on insurance both ways, we can close most of this within a week. We expect D26 to go up when it is priced. That is fine, and it will not be read as your number moving.",
 "Happy to walk through any of it whenever suits.",
 "Matthew",
]
# ---- PDF via Chromium
css="""@page{size:letter;margin:0.9in 1in;} body{font-family:Arial,Helvetica,sans-serif;font-size:10.5pt;color:#1a1a1a;line-height:1.45;}
h1{font-size:16pt;color:#1F3864;margin:0 0 10px 0;} table.h{border-collapse:collapse;margin:0 0 14px 0;} table.h td{padding:1px 10px 1px 0;font-size:9.5pt;}
table.h td.k{color:#595959;font-weight:bold;width:52px;} .rule{border-bottom:1.5px solid #1F3864;margin:0 0 14px 0;} p{margin:0 0 10px 0;}"""
h=[f'<html><head><meta charset="utf-8"><title>{html.escape(TITLE)}</title><style>{css}</style></head><body><h1>{html.escape(TITLE)}</h1><table class="h">']
for k,v in HEAD: h.append(f'<tr><td class="k">{k}</td><td>{html.escape(v)}</td></tr>')
h.append('</table><div class="rule"></div>')
for p in PARAS: h.append(f'<p>{html.escape(p)}</p>')
h.append('</body></html>'); open('cover.html','w',encoding='utf-8').write(''.join(h))
open('topdf.js','w').write("""const { chromium } = require('/tmp/claude-0/-home-user/0521d668-0b2b-59fd-8f23-d788cbd0efd6/scratchpad/node_modules/playwright');
(async()=>{const b=await chromium.launch({executablePath:process.env.PWEXE});const p=await b.newPage();
await p.goto('file://'+process.cwd()+'/cover.html',{waitUntil:'load'});await p.pdf({path:process.argv[2],format:'Letter',printBackground:true,preferCSSPageSize:true});await b.close();})();""")
PDF='165_Randolph_The_BLOCK_Cover_Memo.pdf'
subprocess.run(['node','topdf.js',PDF],check=True,capture_output=True)
print('cover pdf pages:', len(re.findall(rb'/Type\s*/Page[^s]', open(PDF,'rb').read())))
# ---- DOCX
NAVY=RGBColor(0x1F,0x38,0x64); GREY=RGBColor(0x59,0x59,0x59)
d=Document()
for s in d.sections: s.left_margin=s.right_margin=Inches(1); s.top_margin=s.bottom_margin=Inches(0.9)
st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
p=d.add_paragraph(); r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=NAVY
t=d.add_table(rows=0,cols=2)
for k,v in HEAD:
    c=t.add_row().cells; c[0].width=Inches(0.7); c[1].width=Inches(5.8)
    r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    r=c[1].paragraphs[0].add_run(v); r.font.size=Pt(9.5)
    for cc in c: cc.paragraphs[0].paragraph_format.space_after=Pt(0)
d.add_paragraph()
for para in PARAS:
    p=d.add_paragraph(para); p.paragraph_format.space_after=Pt(8)
d.save('165_Randolph_The_BLOCK_Cover_Memo.docx'); print('cover docx ok')
open('cover_memo.txt','w',encoding='utf-8').write('\n\n'.join(PARAS))
