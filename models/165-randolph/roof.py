# -*- coding: utf-8 -*-
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F='Arial'
def build_roof(wb,S):
    BLUE,BLK,BOLD,H1,H2,WHT,SM,RED=S['BLUE'],S['BLK'],S['BOLD'],S['H1'],S['H2'],S['WHT'],S['SM'],S['RED']
    HDR,YEL,BOX,TOPB,CUR,wrap=S['HDR'],S['YEL'],S['BOX'],S['TOPB'],S['CUR'],S['wrap']
    rl=wb.create_sheet('Roof Lift',1)
    for col,w in zip('ABCDEFG',[3,34,12,10,16,4,74]): rl.column_dimensions[col].width=w
    rl['B2']='THE ROOF LIFT'; rl['B2'].font=H1
    rl['B3']='Its own model, because it is the biggest single variable and it is now a DIAL, not a switch.'; rl['B3'].font=SM
    rl['B5']=('With a ground-supported stage and mezzanine (Factory Town style) the roof carries no rigging '
              'load. Height now buys volume, sightlines, LED headroom and HVAC stratification — not tonnage. '
              'So the question stops being "50 ft or nothing" and becomes "how much height does the room need".')
    rl['B5'].font=BLK; rl['B5'].alignment=wrap(); rl.merge_cells('B5:G5'); rl.row_dimensions[5].height=30
    rl['B6']='INPUTS — edit the yellow cells'; rl['B6'].font=H2
    inp=[('Building footprint (SF, under roof)',74100,'#,##0','Your figure. The mezzanine sits INSIDE this — it does not change the envelope.'),
         ('Existing clear height (ft)',20,'#,##0','Per the underwriting: roof raise "~20 -> 50 ft".'),
         ('Target clear height (ft)',50,'#,##0','THE DIAL. Try 35-40: enough for a ground-supported tower system plus LED at 7,440 cap.'),
         ('Lift height (ft)',None,'#,##0','Target less existing.'),
         ('Perimeter (LF)',None,'#,##0','Approximated as a square footprint. Override with the real survey figure.'),
         ('Roof steel (lb/SF)',15,'#,##0','15 with a ground-supported stage. Use 20+ ONLY if the roof must carry rigging.')]
    r=7
    for lab,val,fmt,note in inp:
        rl.cell(r,2,lab).font=BLK
        if val is None:
            f='=E9-E8' if 'Lift height' in lab else '=4*SQRT(E7)'
            c=rl.cell(r,5,f); c.font=BLK
        else:
            c=rl.cell(r,5,val); c.font=BLUE; c.fill=YEL
        c.number_format=fmt; c.border=BOX
        rl.cell(r,7,note).font=SM; rl.cell(r,7).alignment=wrap(); r+=1
    # rows 7..12 = footprint, existing, target, lift, perimeter, lb/SF
    rl['B13']='COST BUILD-UP — at the estimator\'s own unit rates'; rl['B13'].font=H2
    for i,h in enumerate(['Scope','Qty','Unit','Rate','Amount'],2):
        if i==6: continue
    hdr_row=13
    items=[('Demolish existing roof, deck & framing','=E7','SF',15,'=E7*15','File rate $15/SF. Your 74,100 SF footprint.'),
           ('Temporary roof during construction','=E7','SF',5,'=E7*5','File rate $5/SF.'),
           ('New roof trusses, girders & columns','=E7*E12/2000','TON',6000,'=E7*E12/2000*6000','File rate $6,000/ton. NO rigging capacity — the stage stands on the slab.'),
           ('Fireproofing on new roof steel','=E7','SF',5,'=E7*5','File rate $5/SF. NYC-required on new structural steel.'),
           ('New roof membrane','=E7','SF',35,'=E7*35','File rate $35/SF.'),
           ('Masonry "bump up" facade','=E11*E10','SF',35,'=E11*E10*35','File rate $35/SF. Scales directly with the lift height.'),
           ('Furring & sheathing at bump up','=E11*E10','SF',20,'=E11*E10*20','File rate $20/SF. Scales with lift height.'),
           ('New column footings & excavation','40','EA',7500,'=40*7500+E7*6.5','File rates: $7,500/EA footing, $6.50/SF excavation.')]
    r=14
    for nm,qty,unit,rate,amt,note in items:
        rl.cell(r,2,nm).font=BLK
        rl.cell(r,3,qty).number_format='#,##0'
        rl.cell(r,4,unit).alignment=Alignment(horizontal='center')
        c=rl.cell(r,5,amt); c.number_format=CUR; c.font=BLK; c.border=BOX
        rl.cell(r,7,note).font=SM; rl.cell(r,7).alignment=wrap(); r+=1
    rl.cell(r,2,'TRADE SUBTOTAL — the lift').font=BOLD; rl.cell(r,2).border=TOPB
    c=rl.cell(r,5,'=SUM(E14:E21)'); c.font=BOLD; c.number_format=CUR; c.border=TOPB
    rl.cell(r,7,'This total feeds the eight Tier R lines on Line Build. Markups are applied there.').font=SM
    rl.cell(r,7).alignment=wrap()
    r+=2
    rl.cell(r,2,'ROOFLIFTERS 2022 BUDGET RATE CARD  (v1013.1.RL SE)').font=RED; r+=1
    rl.cell(r,2,('A rate card, NOT a project quote — Stage 1 of Rooflifters own three-stage process, priced on '
                 'height and area alone. Rates are "from" figures for NON-UNION projects. At 74,100 SF the card '
                 'interpolates to $20.18/SF between the 50,000 SF ($25) and 75,000 SF ($20) bands. Two adjustments '
                 'are mandatory before use: escalation from 2022, and a New York union / prevailing-wage premium. '
                 'The card INCLUDES lift, structural steel, external metal siding enclosure, bracing, structural '
                 'drawings and GC coordination. It EXCLUDES demolition (except roof separation), hoarding, fire '
                 'separation, and all MEP work — which is more than half of the build-up above.')).font=BLK
    rl.cell(r,2).alignment=wrap(); rl.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7)
    rl.row_dimensions[r].height=30; r+=1
    QR=r
    for lab,val,fmt,note in [
        ('Rate card $/SF at this area',20.18,'$#,##0.00','Interpolated from the 2022 card. Non-union, "from" pricing.'),
        ('Base figure (2022, non-union)','=E7*E'+str(r),CUR,'Area x rate.'),
        ('Escalation 2022 -> construction',1.35,'0.00"x"','Your estimate was 30-40%. Five years at ~6%/yr.'),
        ('NYC union / prevailing wage premium',1.80,'0.00"x"','THE BIG UNKNOWN. Card is explicitly non-union. NYC structural work runs 1.5-2.2x. Get a union-basis figure from them.'),
        ('ADJUSTED ROOFLIFTERS SCOPE','=E'+str(r+1)+'*E'+str(r+2)+'*E'+str(r+3),CUR,'What the lifter would charge here.'),
        ('Excluded scope, still ours','=E14+E15+E17+E18+E21+320000',CUR,'Roof demo, temp roof, fireproofing, membrane, footings, hoarding — all excluded by the card.'),
        ('TOTAL ROOF LIFT PACKAGE','=E'+str(r+4)+'+E'+str(r+5),CUR,'Compare against the build-up subtotal in E22.')]:
        rl.cell(r,2,lab).font=BLK
        c=rl.cell(r,5,val); c.border=BOX; c.number_format=fmt
        c.font=BLUE if not (isinstance(val,str) and val.startswith('=')) else BLK
        if isinstance(val,(int,float)): c.fill=YEL
        rl.cell(r,7,note).font=SM; rl.cell(r,7).alignment=wrap(); r+=1
    TOTR=r-1
    rl.cell(r,2,'Build-up less Rooflifters route').font=BOLD
    c=rl.cell(r,5,f'=E22-E{TOTR}'); c.number_format=CUR; c.font=BOLD; c.border=TOPB
    rl.cell(r,2).border=TOPB
    rl.cell(r,7,'Positive means the Rooflifters route is cheaper. Most of the saving comes from the enclosure being metal siding rather than masonry and furring — not from the steel.').font=SM
    rl.cell(r,7).alignment=wrap()
    r+=2
    rl.cell(r,2,'HEIGHT SENSITIVITY').font=H2; r+=1
    rl.cell(r,2,'Change the target height in E9 and the whole model moves. Roughly:').font=BLK; r+=1
    for i,h in enumerate(['Target height','Lift','Indicative trade cost of the lift'],2):
        c=rl.cell(r,i,h); c.font=WHT; c.fill=HDR; c.alignment=Alignment(horizontal='center',wrap_text=True)
    r+=1
    for tgt in (30,35,40,45,50):
        rl.cell(r,2,f'{tgt} ft').alignment=Alignment(horizontal='center')
        rl.cell(r,3,f'{tgt-20} ft').alignment=Alignment(horizontal='center')
        c=rl.cell(r,4,f'=$E$14+$E$15+$E$16+$E$17+$E$18+$E$21+($E$11*{tgt-20}*35)+($E$11*{tgt-20}*20)')
        c.number_format=CUR; c.font=BLK
        r+=1
    rl.cell(r,2,'Only the facade and furring scale with height — the roof itself costs the same whether you lift it 10 ft or 30 ft.').font=SM
    rl.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7)
    return rl
