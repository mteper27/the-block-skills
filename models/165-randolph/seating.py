# -*- coding: utf-8 -*-
from openpyxl.styles import Alignment
def build_seating(wb,S,pos=4):
    BLUE,BLK,BOLD,H1,H2,WHT,SM,RED=S['BLUE'],S['BLK'],S['BOLD'],S['H1'],S['H2'],S['WHT'],S['SM'],S['RED']
    HDR,YEL,BOX,TOPB,CUR,wrap=S['HDR'],S['YEL'],S['BOX'],S['TOPB'],S['CUR'],S['wrap']
    sh=wb.create_sheet('Theater Seating',pos)
    for col,w in zip('ABCDEFG',[3,40,16,16,16,4,64]): sh.column_dimensions[col].width=w
    sh['B2']='THEATER SEATING — COMPLIANCE vs A BEAUTIFUL THEATER'; sh['B2'].font=H1
    sh['B3']='The seating is not furniture. It is what classifies the room, and the classification carries the licences.'; sh['B3'].font=SM
    sh['B5']=('The room needs a theater use-classification to hold the liquor and cabaret licences, and the technical '
              'drawings have to show seating that supports it. That sets a FLOOR on what must be built. Everything '
              'above that floor is an aesthetic choice, not a licensing one — and the gap is large.')
    sh['B5'].font=BLK; sh['B5'].alignment=wrap(); sh.merge_cells('B5:G5'); sh.row_dimensions[5].height=32
    r=7
    sh.cell(r,2,'INPUTS — confirm the seat count with counsel before relying on any of this').font=H2; r+=1
    A0=r
    for lab,val,fmt,note in [
      ('Seats required for classification',720,'#,##0','PLACEHOLDER. YOUR COUNSEL MUST SET THIS. It drives everything below.'),
      ('Mezzanine seats available',720,'#,##0','657 box chairs + 63 elevated, per FF&E schedule 301.'),
      ('Floor seats in the filed budget',3714,'#,##0','Matches the Wake test fit floor-seated figure at 7 SF/occupant.'),
      ('Removable chair, $/seat',25,'$#,##0','MT H127: "can get cheaper ones for 25 a seat".'),
      ('Filed folding chair, $/seat',125,'$#,##0','As budgeted in FF&E 301.'),
      ('Fixed theatre seat, $/seat',450,'$#,##0','Upholstered, floor-mounted, standard grade. Premium runs $600-900.')]:
        sh.cell(r,2,lab).font=BLK
        c=sh.cell(r,5,val); c.font=BLUE; c.fill=YEL; c.border=BOX; c.number_format=fmt
        sh.cell(r,7,note).font=SM; sh.cell(r,7).alignment=wrap(); r+=1
    REQ,MEZ,FLR,PRM,PFI,PFX=[A0+i for i in range(6)]
    r+=1
    sh.cell(r,2,'THREE WAYS TO DO IT').font=H2; r+=1
    for i,h in enumerate(['','1. COMPLIANCE\nremovable, mezzanine only','2. AS FILED\nfloor folding + mezz boxes','3. BEAUTIFUL THEATER\nfixed, raked, upholstered'],2):
        c=sh.cell(r,i,h); c.font=WHT; c.fill=HDR; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    sh.row_dimensions[r].height=34; r+=1
    ROWS=[('Seats provided',f'=$E${REQ}',f'=$E${FLR}+$E${MEZ}',f'=$E${FLR}+$E${MEZ}','#,##0'),
          ('Chairs / seats',f'=$E${REQ}*$E${PRM}',f'=$E${FLR}*$E${PFI}+$E${MEZ}*300',f'=($E${FLR}+$E${MEZ})*$E${PFX}',CUR),
          ('Raked risers / tiering','=0','=0',f'=($E${FLR}+$E${MEZ})*180',CUR),
          ('Fixed mounting & floor prep','=0','=0',f'=($E${FLR}+$E${MEZ})*45',CUR),
          ('Aisle lighting, row & seat numbering',f'=$E${REQ}*8',f'=($E${FLR}+$E${MEZ})*8',f'=($E${FLR}+$E${MEZ})*22',CUR),
          ('Sightline & seating design fees','=15000','=25000','=140000',CUR),
          ('Storage for removed seating','=35000','=60000','=0',CUR)]
    first=r
    for lab,a,b,c_,fmt in ROWS:
        sh.cell(r,2,lab).font=BLK
        for j,f in enumerate((a,b,c_)):
            cc=sh.cell(r,3+j,f); cc.number_format=fmt; cc.font=BLK; cc.border=BOX
        r+=1
    last=r-1
    sh.cell(r,2,'TOTAL SEATING PACKAGE').font=BOLD; sh.cell(r,2).border=TOPB
    for j,col in enumerate('CDE'):
        cc=sh.cell(r,3+j,f'=SUM({col}{first+1}:{col}{last})'); cc.font=BOLD; cc.number_format=CUR; cc.border=TOPB
    TOT=r; r+=1
    sh.cell(r,2,'Cost per seat provided').font=BLK
    for j,col in enumerate('CDE'):
        cc=sh.cell(r,3+j,f'={col}{TOT}/{col}{first}'); cc.number_format='$#,##0'; cc.font=BLK
    r+=1
    sh.cell(r,2,'Saving vs a beautiful theater').font=BOLD
    for j,col in enumerate('CDE'):
        cc=sh.cell(r,3+j,f'=$E${TOT}-{col}{TOT}'); cc.number_format=CUR; cc.font=BOLD
    r+=2
    sh.cell(r,2,'READ THIS').font=RED; r+=1
    for t in ['Option 1 is the licensing floor: removable seating on the mezzanine, shown on the technical drawings, satisfying the classification — with the floor left as standing room, which is what an Infinity room wants anyway. Seats stack into storage on dance nights.',
              'Option 2 is what is in the budget today. It seats the floor as well as the mezzanine at $125 a chair, and it is neither the cheapest compliant answer nor a real theater.',
              'Option 3 is a genuine theater: fixed upholstered seats on raked tiering with designed sightlines. It is a different building, and on this concept it buys nothing the licence requires.',
              'THE SEAT COUNT IS A PLACEHOLDER. What actually satisfies the NYC theater use-classification, the SLA and the cabaret licence is a legal question: the required count, whether seats must be fixed or may be removable, and whether the mezzanine alone is sufficient. Get that answer before anyone prices a chair.']:
        sh.cell(r,2,'•').font=BLK
        c=sh.cell(r,3,t); c.font=BOLD if 'PLACEHOLDER' in t else BLK; c.alignment=wrap()
        sh.merge_cells(start_row=r,start_column=3,end_row=r,end_column=7)
        sh.row_dimensions[r].height=max(14,12.5*(len(t)//92+1)); r+=1
    return sh
