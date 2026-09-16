# -*- coding: utf-8 -*-
import sys; sys.path.insert(0, '.')
import apples as A, sept16_raw as R
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F = 'Arial'
BLK = Font(name=F, size=9);  BOLD = Font(name=F, size=9, bold=True)
RED = Font(name=F, size=9, bold=True, color='C00000')
GRN = Font(name=F, size=9, bold=True, color='1F6B3B')
H1  = Font(name=F, size=18, bold=True, color='1F3864')
H2  = Font(name=F, size=11, bold=True, color='1F3864')
WHT = Font(name=F, size=9, bold=True, color='FFFFFF')
SM  = Font(name=F, size=8, color='595959')
BIG = Font(name=F, size=14, bold=True, color='1F3864')
SUB = Font(name=F, size=9, bold=True, color='1F3864')
HDR = PatternFill('solid', fgColor='1F3864'); GPF = PatternFill('solid', fgColor='D9E2F3')
PNK = PatternFill('solid', fgColor='FCE4E4'); GRY = PatternFill('solid', fgColor='F2F2F2')
med = Side(style='medium', color='1F3864'); TOPB = Border(top=med)
CUR = '$#,##0;($#,##0);"-"'
def wr(): return Alignment(wrap_text=True, vertical='top')

wb = Workbook()
def head(ws, sub):
    ws.cell(2, 2, '165 RANDOLPH STREET').font = H1
    ws.cell(3, 2, sub).font = H2
    ws.cell(4, 2, 'Our budget and his, complete, in one identical structure · nothing '
                  'stripped from either side · 16 Sep 2026 · CONFIDENTIAL').font = SM
    return 6
def hrow(ws, r, cols):
    for i, h in enumerate(cols, 2):
        c = ws.cell(r, i, h); c.font = WHT; c.fill = HDR
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    return r + 1

# ================================================== 1. APPLES TO APPLES
ws = wb.active; ws.title = 'Apples to Apples'
for col, w in zip('ABCDEF', [3, 50, 15, 15, 15, 96]): ws.column_dimensions[col].width = w
r = head(ws, 'OUR BUDGET AND HIS, LINE FOR LINE')
ws.cell(r, 2, 'Ours $%s.  His $%s.  He is $%s higher.'
        % (f'{A.OURS_TOTAL:,.0f}', f'{R.TOTAL:,.0f}',
           f'{R.TOTAL - A.OURS_TOTAL:,.0f}')).font = BIG
r += 2
for t in [
  'Every row means the same thing in both columns, and each column foots to the total that '
  'budget actually publishes. Two normalisations, and only two, so the rows line up — '
  'neither changes either total:',
  '     1.  His division 013000 PROJECT REQUIREMENTS ($3,034,000) and 001000 PRECONSTRUCTION '
  '($60,000) move out of his trade cost onto the running-the-site rows, where our equivalent '
  'already sits.',
  '     2.  His abatement ($1,483,680) moves out of his soft costs into the hazmat work '
  'package, where ours already sits.']:
    c = ws.cell(r, 2, t); c.font = BLK; c.alignment = wr()
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 26 if len(t) > 150 else 13
    r += 1
r += 1
r = hrow(ws, r, ['', 'OURS', 'HIS', 'Delta', 'What is going on'])
for kind, lab, o, h, note in A.ROWS:
    ind = '     ' if kind == 'pkg' else ''
    c = ws.cell(r, 2, ind + lab)
    c.font = BLK if kind == 'pkg' else (BOLD if kind == 'tot' else SUB)
    for col, v in ((3, o), (4, h), (5, h - o)):
        cc = ws.cell(r, col, v); cc.number_format = CUR
        if col == 5 and kind == 'pkg':
            cc.font = RED if h > o else (GRN if h < o else BLK)
        else:
            cc.font = BLK if kind == 'pkg' else (BOLD if kind == 'tot' else SUB)
    if kind == 'pkg' and h == 0: ws.cell(r, 4).fill = PNK
    if kind in ('sub', 'tot'):
        fill = GPF if kind == 'tot' else GRY
        for col in range(2, 7):
            ws.cell(r, col).fill = fill; ws.cell(r, col).border = TOPB
    c = ws.cell(r, 6, note.replace('%%', '%')); c.font = SM if kind == 'pkg' else BLK
    c.alignment = wr()
    ws.row_dimensions[r].height = max(14 if kind == 'pkg' else 16,
                                      10.5 * (len(note) // 98 + 1))
    r += 1
r += 1
ws.cell(r, 2, 'THE SAME TWO BUDGETS, FIVE WAYS').font = H2; r += 1
r = hrow(ws, r, ['', 'OURS', 'HIS', 'Delta'])
for lab, o, h, fmt in A.METRICS:
    ws.cell(r, 2, lab.replace('%%', '%')).font = BLK
    for col, v in ((3, o), (4, h)):
        c = ws.cell(r, col, v); c.number_format = fmt; c.font = BOLD
    c = ws.cell(r, 5, h - o); c.number_format = fmt; c.font = RED
    r += 1
r += 1
ws.cell(r, 2, 'His own sheet reports $655/SF — it divides by 135,953 SF, which is 555 '
              'Johnson, not this building.').font = SM

# ========================================================== 2. THE ISSUES
ws = wb.create_sheet('The Issues')
for col, w in zip('ABCDEF', [3, 10, 50, 74, 74, 16]): ws.column_dimensions[col].width = w
r = head(ws, 'WHAT ACCOUNTS FOR THE DIFFERENCE')
ws.cell(r, 2, 'HIS = his number needs to move.   OURS = ours does.   BOTH = a decision neither '
              'of us can make alone.   NEITHER = leave it exactly where it is.').font = SM
r += 2
r = hrow(ws, r, ['Whose', 'Issue', 'The evidence, quoted from his file', 'What to do',
                 'All-in $ at stake, indicative'])
for whose, title, ev, act, amt in A.ISSUES:
    c = ws.cell(r, 2, whose)
    c.font = RED if whose == 'HIS' else (GRN if whose == 'OURS' else BOLD)
    c.alignment = Alignment(horizontal='center', vertical='top')
    c = ws.cell(r, 3, title); c.font = BOLD; c.alignment = wr()
    c = ws.cell(r, 4, ev); c.font = BLK; c.alignment = wr()
    c = ws.cell(r, 5, act); c.font = BLK; c.alignment = wr()
    c = ws.cell(r, 6, amt if amt else None); c.number_format = CUR; c.font = BOLD
    ws.row_dimensions[r].height = max(46, 10.5 * (max(len(ev), len(act)) // 76 + 1))
    r += 1
r += 1
quant = sum(i[4] for i in A.ISSUES)
ws.cell(r, 3, 'QUANTIFIED, ALL-IN').font = BOLD
c = ws.cell(r, 6, quant); c.number_format = CUR; c.font = BOLD; c.border = TOPB
ws.cell(r, 3).border = TOPB; r += 2
for t in ['The seven quantified items come to $%s of his $%s gap. Each is priced on its own, '
          'so they are INDICATIVE and do not sum exactly — correcting one changes the base '
          'the next one sits on.' % (f'{quant:,.0f}', f'{R.TOTAL - A.OURS_TOTAL:,.0f}'),
          'For the exact combined effect, the audit workbook applies every correction to a '
          'source cell and re-runs his own formula chain: $89,084,591 falls to $68,143,696, a '
          'real reduction of $20,940,896.',
          'What is left is programme — a second floor, roof decks, a balcony, two lounges, '
          'a foundry and a restaurant that he priced because the test fit showed them to him. '
          'That is not an estimating error and it is not his to fix.']:
    c = ws.cell(r, 3, t); c.font = BLK; c.alignment = wr()
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 26; r += 1

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.freeze_panes = 'A7'
wb.properties.title = '165 Randolph - Apples to Apples'
wb.calculation.fullCalcOnLoad = True
OUT = '165_Randolph_Apples_to_Apples_Ours_vs_His.xlsx'
wb.save(OUT)
print('wrote', OUT)
